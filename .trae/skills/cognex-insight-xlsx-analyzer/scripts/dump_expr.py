import openpyxl, sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
ws = wb['单元格']
start, end = int(sys.argv[2]), int(sys.argv[3])
seen = set()
for r in range(start, end+1):
    addr = ws.cell(r,1).value
    val = ws.cell(r,3).value
    expr = ws.cell(r,4).value
    if not isinstance(expr, str) or len(expr) < 2:
        continue
    if not isinstance(addr, str) or not re.match(r'^[A-Z]+\d+$', addr):
        continue
    key = addr + '|' + str(expr)
    if key in seen:
        continue
    seen.add(key)
    v = str(val).replace('\n',' ') if val is not None else ''
    if len(v) > 90: v = v[:90]
    print(addr, '=', v, ' <- ', expr)
