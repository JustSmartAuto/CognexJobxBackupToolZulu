import openpyxl, sys, re, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
ws = wb['单元格']
pat = re.compile(r'^"(\d+[、.][^"]+)"$')
for r in range(1, ws.max_row+1):
    v = ws.cell(r, 3).value
    if isinstance(v, str) and pat.match(v):
        print(r, ws.cell(r, 1).value, pat.match(v).group(1))
